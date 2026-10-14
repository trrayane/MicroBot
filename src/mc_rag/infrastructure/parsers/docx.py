"""DOCX document parser built on top of :mod:`docx` (python-docx).

The public entry point is :func:`parse_docx`. Body paragraphs are rendered in
document order with headings preserved as Markdown-style ``#`` lines, and each
heading starts a :class:`~mc_rag.infrastructure.parsers.base.Section` whose body
is populated by the paragraphs that follow it. Tables are rendered row by row as
``header : value`` lines so tabular data is neither dropped nor flattened into
an unreadable ``a | b | c`` line.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

from mc_rag.domain.exceptions import (
    CorruptDocumentError,
    EmptyDocumentError,
)
from mc_rag.infrastructure.parsers.base import (
    ParsedDocument,
    Section,
    append_to_section,
    coerce_optional_str,
    resolve_document_path,
    slugify,
)
from mc_rag.infrastructure.parsers.table import TABLE_LABEL, build_rows

logger = logging.getLogger(__name__)

_FORMAT = "docx"
_HEADING_STYLE_PREFIX = "heading"
_TITLE_STYLES = frozenset({"title", "document title"})


def parse_docx(file_path: str | Path) -> ParsedDocument:
    """Extract text and structure from a Word ``.docx`` file.

    Args:
        file_path: Path to the DOCX document.

    Returns:
        A :class:`ParsedDocument` whose ``text`` is the flattened body, whose
        ``sections`` mirror heading paragraphs, and whose ``metadata`` records
        the paragraph count, heading count and author (when present).

    Raises:
        DocumentNotFoundError: The path does not exist or is not a file.
        EmptyDocumentError: The file is zero bytes or holds no readable text.
        CorruptDocumentError: The file is not a valid DOCX package.
    """
    path = resolve_document_path(file_path)

    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001 - normalize third-party errors
        logger.warning("Failed to open DOCX %s: %s", path, exc)
        raise CorruptDocumentError(f"Unable to open DOCX: {path}") from exc

    lines, sections, paragraph_count, style_title = _extract_paragraphs(document)
    lines.extend(_extract_tables(document, file_name=path.name))

    text = "\n".join(lines).strip()
    if not text:
        raise EmptyDocumentError(f"DOCX contains no extractable text: {path}")

    core = document.core_properties
    metadata: dict[str, Any] = {
        "paragraph_count": paragraph_count,
        "heading_count": len(sections),
    }
    author = coerce_optional_str(getattr(core, "author", None))
    if author:
        metadata["author"] = author

    title = (
        coerce_optional_str(getattr(core, "title", None))
        or style_title
        or (sections[0].heading if sections else None)
    )

    return ParsedDocument(
        text=text,
        source=str(path),
        format=_FORMAT,
        title=title,
        metadata=metadata,
        sections=sections,
    )


def _extract_paragraphs(
    document: DocxDocument,
) -> tuple[list[str], list[Section], int, str | None]:
    """Render paragraphs to lines, deriving sections from heading styles."""
    lines: list[str] = []
    sections: list[Section] = []
    current: Section | None = None
    paragraph_count = 0
    style_title: str | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph_count += 1

        level = _heading_level(paragraph)
        if level is None:
            lines.append(text)
            current = append_to_section(current, text)
            continue

        if style_title is None and _is_title_style(paragraph):
            style_title = text
        lines.append(f"{'#' * level} {text}")
        current = Section(level=level, heading=text, anchor=slugify(text))
        sections.append(current)

    return lines, sections, paragraph_count, style_title


def _extract_tables(document: DocxDocument, *, file_name: str) -> list[str]:
    """Render table rows as ``header : value`` lines (never as ``a | b | c``).

    The first row of a table is treated as its header when it holds only text;
    otherwise columns are labelled ``Colonne N``. Rows keep the file name and
    the table index so they remain understandable out of context.
    """
    lines: list[str] = []
    for index, table in enumerate(document.tables, start=1):
        grid = [[cell.text for cell in row.cells] for row in table.rows]
        try:
            rows = build_rows(
                grid,
                file_name=file_name,
                container=str(index),
                container_label=TABLE_LABEL,
            )
        except EmptyDocumentError:
            logger.debug("Skipping empty DOCX table %d", index)
            continue
        lines.extend(row.text for row in rows)
    return lines


def _style_name(paragraph: Paragraph) -> str:
    """Return the lower-cased style name of *paragraph*, or ``""``."""
    style = paragraph.style
    name = getattr(style, "name", None)
    return name.strip().lower() if name else ""


def _heading_level(paragraph: Paragraph) -> int | None:
    """Return the heading level for *paragraph*, or ``None`` if not a heading."""
    name = _style_name(paragraph)
    if name not in _TITLE_STYLES and not name.startswith(_HEADING_STYLE_PREFIX):
        return None
    digits = "".join(character for character in name if character.isdigit())
    return max(1, int(digits)) if digits else 1


def _is_title_style(paragraph: Paragraph) -> bool:
    """Whether *paragraph* uses an explicit title style."""
    return _style_name(paragraph) in _TITLE_STYLES
