"""Unit tests for the tabular parsers (Excel, CSV, DOCX tables, Google Sheets).

These tests pin the contract the chunker and the retriever rely on:

* one rendered line per data row, ``header : value`` pairs, never ``a | b | c``;
* file name (and sheet/table name) inside every row;
* ``text == "\\n".join(row.text for row in tabular)``;
* explicit errors for empty, corrupt and oversized spreadsheets.
"""

from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from mc_rag.domain.exceptions import (
    CorruptDocumentError,
    DocumentTooLargeError,
    EmptyDocumentError,
    ParsingError,
)
from mc_rag.infrastructure.parsers.csv import parse_csv, parse_csv_text
from mc_rag.infrastructure.parsers.docx import parse_docx
from mc_rag.infrastructure.parsers.gsheets import (
    FetchResult,
    is_google_sheets_url,
    parse_google_sheets,
)
from mc_rag.infrastructure.parsers.table import (
    build_rows,
    format_cell,
    looks_like_header,
)
from mc_rag.infrastructure.parsers.xlsx import parse_xlsx, parse_xlsx_bytes

SHEET_URL = "https://docs.google.com/spreadsheets/d/abcDEFGH1234/edit#gid=0"


def write_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    """Write an ``.xlsx`` file holding *sheets* (order preserved)."""
    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = name
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return path


def workbook_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    """Return the raw bytes of an in-memory workbook holding *sheets*."""
    buffer = io.BytesIO()
    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = name
        for row in rows:
            sheet.append(row)
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# build_rows: shared rendering rules
# --------------------------------------------------------------------------- #


def test_build_rows_renders_file_sheet_and_pairs() -> None:
    rows = build_rows(
        [["Nom", "Rôle"], ["Amine", "Trésorier"], ["Sara", "Secrétaire"]],
        file_name="bureau.xlsx",
        container="Bureau",
        container_label="Feuille",
        sheet="Bureau",
    )
    assert [row.text for row in rows] == [
        "Fichier : bureau.xlsx | Feuille : Bureau | Nom : Amine | Rôle : Trésorier",
        "Fichier : bureau.xlsx | Feuille : Bureau | Nom : Sara | Rôle : Secrétaire",
    ]
    assert rows[0].sheet == "Bureau"
    assert rows[0].row_number == 2
    assert rows[1].row_number == 3


def test_build_rows_without_header_generates_column_labels() -> None:
    rows = build_rows(
        [[1, "Amine"], [2, "Sara"]],
        file_name="bureau.csv",
        container="",
        container_label=None,
    )
    assert rows[0].text == "Fichier : bureau.csv | Colonne 1 : 1 | Colonne 2 : Amine"


def test_build_rows_drops_empty_rows_and_columns() -> None:
    rows = build_rows(
        [["Nom", "", "Rôle"], ["", "", ""], ["Amine", "", "Trésorier"]],
        file_name="bureau.csv",
        container="",
        container_label=None,
    )
    assert [row.text for row in rows] == [
        "Fichier : bureau.csv | Nom : Amine | Rôle : Trésorier",
    ]


def test_build_rows_deduplicates_identical_headers() -> None:
    rows = build_rows(
        [["Nom", "Nom"], ["Amine", "Sara"]],
        file_name="b.csv",
        container="",
        container_label=None,
    )
    assert "Nom : Amine" in rows[0].text
    assert "Nom (2) : Sara" in rows[0].text


def test_build_rows_without_container_has_no_container_label() -> None:
    rows = build_rows(
        [["Nom"], ["Amine"]],
        file_name="b.csv",
        container="",
        container_label=None,
    )
    assert rows[0].text == "Fichier : b.csv | Nom : Amine"


def test_build_rows_without_any_value_raises() -> None:
    with pytest.raises(EmptyDocumentError):
        build_rows(
            [[" ", None], ["", ""]],
            file_name="empty.csv",
            container="",
            container_label=None,
        )


def test_build_rows_enforces_row_budget() -> None:
    with pytest.raises(DocumentTooLargeError, match="SHEET_MAX_ROWS"):
        build_rows(
            [["Nom"], ["A"], ["B"], ["C"]],
            file_name="big.csv",
            container="",
            container_label=None,
            max_rows=2,
        )


def test_build_rows_tolerates_ragged_rows() -> None:
    rows = build_rows(
        [["Nom", "Rôle", "Ville"], ["Amine"]],
        file_name="b.csv",
        container="",
        container_label=None,
    )
    assert rows[0].text == "Fichier : b.csv | Nom : Amine"


def test_build_rows_skips_title_banner_and_finds_the_real_header() -> None:
    """Regression: student lists start with a banner, not with the header.

    The header must be the widest all-text row near the top, otherwise every
    column gets a generated ``Colonne N`` label and student data (matricule,
    groupe TD…) becomes unsearchable.
    """
    rows = build_rows(
        [
            ["Université des Sciences et de Technologie", None, None],
            ["Faculté d'Informatique", None, None],
            ["Liste des étudiants: M1 IV A", None, None],
            [None, None, None],
            ["N°", "Nom", "Groupe TD"],
            [1, "TERKI", 2],
            [2, "AFRIT", 1],
        ],
        file_name="liste.xlsx",
        container="M1 IV A",
        container_label="Feuille",
        sheet="M1 IV A",
    )
    assert [row.text for row in rows] == [
        "Fichier : liste.xlsx | Feuille : M1 IV A | N° : 1 | Nom : TERKI | Groupe TD : 2",
        "Fichier : liste.xlsx | Feuille : M1 IV A | N° : 2 | Nom : AFRIT | Groupe TD : 1",
    ]


def test_build_rows_header_tie_prefers_the_topmost_row() -> None:
    rows = build_rows(
        [["Nom", "Rôle"], ["Amine", "Trésorier"], ["Nom", "Rôle"]],
        file_name="b.csv",
        container="",
        container_label=None,
    )
    assert [row.text for row in rows] == [
        "Fichier : b.csv | Nom : Amine | Rôle : Trésorier",
        "Fichier : b.csv | Nom : Nom | Rôle : Rôle",
    ]


def test_looks_like_header_rules() -> None:
    assert looks_like_header(["Nom", "Rôle"])
    assert not looks_like_header(["Nom", 3])
    assert not looks_like_header([None, ""])


def test_format_cell_is_readable() -> None:
    assert format_cell(None) == ""
    assert format_cell(True) == "True"
    assert format_cell(2.0) == "2"
    assert format_cell(2.5) == "2.5"
    assert format_cell("  a   b  ") == "a b"
    assert format_cell(dt.date(2024, 5, 1)) == "2024-05-01"
    assert format_cell(dt.datetime(2024, 5, 1, 14, 30)) == "2024-05-01 14:30"
    assert format_cell(dt.datetime(2024, 5, 1)) == "2024-05-01"
    assert format_cell("x" * 400, max_chars=10) == "x" * 9 + "…"


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #


def test_parse_xlsx_reads_every_sheet_and_keeps_the_invariant(tmp_path: Path) -> None:
    path = write_workbook(
        tmp_path / "bureau.xlsx",
        {
            "Bureau": [["Nom", "Rôle"], ["Amine", "Trésorier"]],
            "Vides": [[" ", None]],
            "Contacts": [["Nom", "Téléphone"], ["Sara", "0612345678"]],
        },
    )
    document = parse_xlsx(path)

    assert document.format == "xlsx"
    assert document.source == str(path)
    assert document.title == "bureau"
    assert document.metadata["sheets"] == ["Bureau", "Contacts"]
    assert document.metadata["sheet_count"] == 2
    assert document.tabular is not None
    assert len(document.tabular) == 2
    assert document.text == "\n".join(row.text for row in document.tabular)
    assert "Fichier : bureau.xlsx | Feuille : Bureau | Nom : Amine" in document.text
    assert "Feuille : Contacts | Nom : Sara | Téléphone : 0612345678" in document.text


def test_parse_xlsx_reports_cell_types(tmp_path: Path) -> None:
    path = write_workbook(
        tmp_path / "suivi.xlsx",
        {
            "Suivi": [
                ["Agent", "Heures", "Jour"],
                ["Amine", 8, dt.datetime(2024, 5, 1)],
            ]
        },
    )
    document = parse_xlsx(path)
    row = (document.tabular or [])[0].text
    assert "Agent : Amine" in row
    assert "Heures : 8" in row
    assert "Jour : 2024-05-01" in row


def test_parse_xlsx_vertical_merge_repeats_the_value(tmp_path: Path) -> None:
    path = write_workbook(
        tmp_path / "merged.xlsx",
        {"F": [["Ville"], ["Paris"], [None], ["Lyon"]]},
    )
    workbook = load_workbook(path)
    workbook["F"].merge_cells("A2:A3")
    workbook.save(path)

    document = parse_xlsx(path)
    rendered = [row.text for row in document.tabular or []]
    assert rendered == [
        "Fichier : merged.xlsx | Feuille : F | Ville : Paris",
        "Fichier : merged.xlsx | Feuille : F | Ville : Paris",
        "Fichier : merged.xlsx | Feuille : F | Ville : Lyon",
    ]


def test_parse_xlsx_horizontal_merge_does_not_duplicate(tmp_path: Path) -> None:
    path = write_workbook(
        tmp_path / "hmerge.xlsx",
        {"F": [["Nom", "Adresse"], ["Amine", "Paris"]]},
    )
    workbook = load_workbook(path)
    workbook["F"].merge_cells("B2:C2")
    workbook.save(path)

    document = parse_xlsx(path)
    assert [row.text for row in document.tabular or []] == [
        "Fichier : hmerge.xlsx | Feuille : F | Nom : Amine | Adresse : Paris",
    ]


def test_parse_xlsx_rejects_corrupt_payload(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"this is not a zip archive")
    with pytest.raises(CorruptDocumentError):
        parse_xlsx(path)


def test_parse_xlsx_rejects_empty_file(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    path.write_bytes(b"")
    with pytest.raises(EmptyDocumentError):
        parse_xlsx(path)


def test_parse_xlsx_enforces_row_budget(tmp_path: Path) -> None:
    path = write_workbook(
        tmp_path / "big.xlsx", {"F": [["Nom"], ["A"], ["B"], ["C"]]}
    )
    with pytest.raises(DocumentTooLargeError):
        parse_xlsx(path, max_rows=2)


def test_parse_xlsx_bytes_requires_payload() -> None:
    with pytest.raises(EmptyDocumentError):
        parse_xlsx_bytes(b"", source="<memory>")


def test_parse_xlsx_formula_without_cached_value_reads_as_empty(tmp_path: Path) -> None:
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "F"
    sheet["A1"] = "Total"
    sheet["B1"] = "=1+1"
    sheet["A2"] = "Année"
    sheet["B2"] = "=2024*2"
    workbook.save(path)

    # openpyxl writes formulas without a cached result: without Excel opening
    # the file first, the value is unknown and reads as empty (documented).
    document = parse_xlsx(path)
    assert document.tabular is not None
    assert all("=1+1" not in row.text for row in document.tabular)


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #


def test_parse_csv_semicolon_with_bom(tmp_path: Path) -> None:
    path = tmp_path / "bureau.csv"
    path.write_bytes("﻿Nom;Rôle\nAmine;Trésorier\n".encode())

    document = parse_csv(path)
    assert document.format == "csv"
    assert document.metadata["delimiter"] == ";"
    assert document.tabular is not None
    assert [row.text for row in document.tabular] == [
        "Fichier : bureau.csv | Nom : Amine | Rôle : Trésorier",
    ]
    assert document.text == "\n".join(row.text for row in document.tabular)


def test_parse_csv_comma(tmp_path: Path) -> None:
    path = tmp_path / "bureau.csv"
    path.write_text("Nom,Rôle\nAmine,Trésorier\n", encoding="utf-8")
    document = parse_csv(path)
    assert document.metadata["delimiter"] == ","
    assert "Nom : Amine" in document.text


def test_parse_csv_legacy_encoding(tmp_path: Path) -> None:
    path = tmp_path / "latin.csv"
    path.write_bytes("Nom;Rôle\nAmine;Trésorier\n".encode("cp1252"))
    document = parse_csv(path)
    assert "Rôle : Trésorier" in document.text


def test_parse_csv_header_only_is_empty(tmp_path: Path) -> None:
    path = tmp_path / "header.csv"
    path.write_text("Nom;Rôle\n", encoding="utf-8")
    with pytest.raises(EmptyDocumentError):
        parse_csv(path)


def test_parse_csv_empty_file(tmp_path: Path) -> None:
    path = tmp_path / "empty.csv"
    path.write_bytes(b"")
    with pytest.raises(EmptyDocumentError):
        parse_csv(path)


def test_parse_csv_text_without_delimiter_raises() -> None:
    with pytest.raises(ParsingError, match="delimiter"):
        parse_csv_text("une seule colonne libre", source="<memory>")


# --------------------------------------------------------------------------- #
# DOCX tables
# --------------------------------------------------------------------------- #


def test_docx_table_rows_render_header_value_lines(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Rapport mensuel")
    table = document.add_table(rows=1, cols=2)
    header = table.rows[0].cells
    header[0].text = "Nom"
    header[1].text = "Rôle"
    cells = table.add_row().cells
    cells[0].text = "Amine"
    cells[1].text = "Trésorier"
    path = tmp_path / "rapport.docx"
    document.save(path)

    parsed = parse_docx(path)
    lines = parsed.text.splitlines()
    assert any("Nom : Amine" in line for line in lines)
    assert any("Rôle : Trésorier" in line for line in lines)
    assert all("Amine | Trésorier" not in line for line in lines)
    assert all("rapport.docx" in line for line in lines if "Nom :" in line)
    assert parsed.tabular is None


# --------------------------------------------------------------------------- #
# Google Sheets
# --------------------------------------------------------------------------- #


def test_is_google_sheets_url_detects_edit_and_export_links() -> None:
    assert is_google_sheets_url(SHEET_URL)
    assert is_google_sheets_url(
        "https://docs.google.com/spreadsheets/d/abcDEFGH1234/export?format=csv"
    )
    assert is_google_sheets_url("https://docs.google.com/spreadsheets/u/0/d/xyz1234/edit")
    assert not is_google_sheets_url("https://example.com/spreadsheets/d/abc/edit")
    assert not is_google_sheets_url("https://docs.google.com/document/d/abc/edit")


def test_google_sheets_rejects_non_sheet_url() -> None:
    with pytest.raises(ParsingError, match="Not a Google Sheets URL"):
        parse_google_sheets("https://example.com/sheet", fetcher=lambda _url: _missing())


def test_google_sheets_private_sheet_gives_actionable_error() -> None:
    def fetcher(_url: str) -> FetchResult:
        return FetchResult(
            url="https://accounts.google.com/ServiceLogin?continue=x",
            status_code=200,
            content_type="text/html",
            body=b"<html><body>Sign in</body></html>",
        )

    with pytest.raises(ParsingError, match="not shared publicly"):
        parse_google_sheets(SHEET_URL, fetcher=fetcher)


def test_google_sheets_parses_the_xlsx_export() -> None:
    payload = workbook_bytes(
        {"Bureau": [["Nom", "Rôle"], ["Amine", "Trésorier"]]}
    )
    seen: list[str] = []

    def fetcher(url: str) -> FetchResult:
        seen.append(url)
        return FetchResult(
            url=url, status_code=200, content_type="application/octet-stream", body=payload
        )

    document = parse_google_sheets(SHEET_URL, fetcher=fetcher)
    assert seen == ["https://docs.google.com/spreadsheets/d/abcDEFGH1234/export?format=xlsx"]
    assert document.format == "gsheet"
    assert document.source == SHEET_URL
    assert document.title == "Google Sheet"
    assert document.tabular is not None
    assert document.tabular[0].text == (
        "Fichier : Google Sheet | Feuille : Bureau | Nom : Amine | Rôle : Trésorier"
    )


def test_google_sheets_falls_back_to_csv_export() -> None:
    def fetcher(url: str) -> FetchResult:
        if "format=xlsx" in url:
            return FetchResult(
                url=url, status_code=500, content_type="text/html", body=b"<html/>"
            )
        return FetchResult(
            url=url,
            status_code=200,
            content_type="text/csv",
            body="Nom;Rôle\nAmine;Trésorier\n".encode(),
        )

    document = parse_google_sheets(SHEET_URL, fetcher=fetcher)
    assert document.format == "gsheet"
    assert document.tabular is not None
    assert document.tabular[0].text == (
        "Fichier : Google Sheet | Nom : Amine | Rôle : Trésorier"
    )


def test_google_sheets_raises_when_every_export_fails() -> None:
    def fetcher(url: str) -> FetchResult:
        return FetchResult(url=url, status_code=500, content_type="text/html", body=b"boom")

    with pytest.raises(ParsingError, match="Unable to export"):
        parse_google_sheets(SHEET_URL, fetcher=fetcher)


def _missing() -> FetchResult:
    """Placeholder result used where no fetch should happen."""
    return FetchResult(url="", status_code=500, content_type="", body=b"")
